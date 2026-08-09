#!/usr/bin/env python3
"""게릴라 배포 기록표 — 구글 시트 import용 xlsx.

설계 원칙: 사장님이 손으로 적는 건 '배포 대수'와 '문의 1건'뿐이다.
문의 수·전환 수·전환율·단지 판정은 전부 수식이 계산한다.
현장에서 계산하게 만들면 기록이 끊긴다.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

F = "Arial"
NAVY = "16365F"
ACCENT = "0F7FBD"

ttl = Font(name=F, size=14, bold=True, color=NAVY)
hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
bod = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
inp = Font(name=F, size=10, color="0000FF")          # 파랑 = 직접 입력
note = Font(name=F, size=9, color="6B7A8C", italic=True)
big = Font(name=F, size=12, bold=True, color=NAVY)

hdr_fill = PatternFill("solid", fgColor=NAVY)
in_fill = PatternFill("solid", fgColor="FFFF99")     # 노랑 = 입력 칸
ex_fill = PatternFill("solid", fgColor="EFF3F7")     # 예시 행
sum_fill = PatternFill("solid", fgColor="E3F0F9")

thin = Side(style="thin", color="D7DEE7")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center")

PCT = "0.0%"
NUM = "#,##0"
WON = '#,##0"원"'
DATE = "yyyy-mm-dd"

BLOCKS = ["A", "B", "C"]
ROUNDS = ["1", "2", "3"]
CODES = [f"WJ-{b}{r}" for b in BLOCKS for r in ROUNDS]

LOG_TOP, LOG_BOT = 4, 503          # 문의접수 데이터 영역
DIST_TOP = 4                        # 배포기록 첫 행
DIST_BOT = DIST_TOP + len(CODES) - 1
DIST_SUM = DIST_BOT + 1

wb = Workbook()

# ─────────────────────────────────────────────────────────────
# 1. 요약
# ─────────────────────────────────────────────────────────────
s = wb.active
s.title = "요약"
s.sheet_view.showGridLines = False

s["A1"] = "게릴라 배포 실험 — 요약"
s["A1"].font = ttl
s["A2"] = "플레이북 §7.3 단지 실험. 아래 숫자는 전부 자동 계산됩니다 — 노란 칸만 직접 입력하세요."
s["A2"].font = note

s["A4"] = "판정 기준 전환율"
s["A4"].font = bold
s["B4"] = 0.03
s["B4"].font = inp
s["B4"].fill = in_fill
s["B4"].number_format = PCT
s["B4"].border = box
s["C4"] = "← 플레이북 §7.2 보수 시나리오 하한. 이 값을 넘으면 '집중', 못 넘으면 '보류'로 판정합니다."
s["C4"].font = note


def block_table(ws, row, title, labels, key_col, code_prefix_len):
    """단지별 / 회차별 집계 블록을 그린다."""
    ws.cell(row, 1, title).font = big
    heads = ["구분", "배포 대수", "문의 수", "구독 전환", "문의율", "전환율", "판정"]
    r = row + 1
    for i, h in enumerate(heads, start=1):
        c = ws.cell(r, i, h)
        c.font, c.fill, c.alignment, c.border = hdr, hdr_fill, ctr, box
    for j, lab in enumerate(labels):
        rr = r + 1 + j
        key = lab[0]
        ws.cell(rr, 1, lab).font = bod
        ws.cell(rr, 2, f'=SUMIF(배포기록!${key_col}${DIST_TOP}:${key_col}${DIST_BOT},"{key}",배포기록!$D${DIST_TOP}:$D${DIST_BOT})')
        ws.cell(rr, 3, f'=SUMIF(배포기록!${key_col}${DIST_TOP}:${key_col}${DIST_BOT},"{key}",배포기록!$E${DIST_TOP}:$E${DIST_BOT})')
        ws.cell(rr, 4, f'=SUMIF(배포기록!${key_col}${DIST_TOP}:${key_col}${DIST_BOT},"{key}",배포기록!$F${DIST_TOP}:$F${DIST_BOT})')
        ws.cell(rr, 5, f'=IFERROR(C{rr}/B{rr},"")')
        ws.cell(rr, 6, f'=IFERROR(D{rr}/B{rr},"")')
        ws.cell(rr, 7, f'=IF(B{rr}=0,"미집계",IF(F{rr}>=요약!$B$4,"집중","보류"))')
        for i in range(1, 8):
            cc = ws.cell(rr, i)
            cc.border = box
            if cc.font.name is None or i > 1:
                cc.font = bod
            if i in (2, 3, 4):
                cc.number_format = NUM
            if i in (5, 6):
                cc.number_format = PCT
            if i == 7:
                cc.alignment = ctr
    return r + len(labels)


last = block_table(s, 6, "단지별 — 어디에 자원을 몰 것인가",
                   ["A 단지", "B 단지", "C 단지"], "J", 1)

r2 = block_table(s, last + 3, "회차별 — 반복 노출이 전환을 올리는가",
                 ["1차 배포", "2차 배포", "3차 배포"], "K", 1)

# 종합
r = r2 + 3
s.cell(r, 1, "종합").font = big
labs = [
    ("총 배포 대수", f"=배포기록!D{DIST_SUM}", NUM),
    ("총 문의 수", f"=배포기록!E{DIST_SUM}", NUM),
    ("총 구독 전환", f"=배포기록!F{DIST_SUM}", NUM),
    ("종합 전환율", f"=IFERROR(B{r+3}/B{r+1},\"\")", PCT),
    ("월 증가 매출", f"=B{r+3}*80000", WON),
]
for i, (lab, fx, fmt) in enumerate(labs):
    rr = r + 1 + i
    s.cell(rr, 1, lab).font = bold if i == 3 else bod
    c = s.cell(rr, 2, fx)
    c.font = bold if i == 3 else bod
    c.number_format = fmt
    c.fill = sum_fill
    c.border = box
    s.cell(rr, 1).border = box

rr = r + 6
s.cell(rr, 1, "자원 집중 대상").font = bold
s.cell(rr, 2, f'=IF(SUM(B{r+1})=0,"데이터 없음",'
              f'INDEX($A$8:$A$10,MATCH(MAX($F$8:$F$10),$F$8:$F$10,0)))')
s.cell(rr, 2).font = Font(name=F, size=11, bold=True, color="0F7FBD")
s.cell(rr, 2).fill = sum_fill
s.cell(rr, 2).border = box
s.cell(rr, 1).border = box
s.cell(rr, 3, "← 4주 후 이 단지에 자원을 몰아넣고 나머지에서 뺍니다 (플레이북 §0 초집중 시나리오).").font = note

for col, w in zip("ABCDEFG", [18, 14, 12, 12, 11, 11, 11]):
    s.column_dimensions[col].width = w
s.column_dimensions["C"].width = 14

# ─────────────────────────────────────────────────────────────
# 2. 배포기록
# ─────────────────────────────────────────────────────────────
d = wb.create_sheet("배포기록")
d.sheet_view.showGridLines = False
d["A1"] = "배포 기록 — 코드별 분모"
d["A1"].font = ttl
d["A2"] = ("노란 칸(단지명·배포일·배포 대수)만 직접 입력하세요. "
           "문의 수부터는 '문의접수' 시트에서 자동으로 집계됩니다.")
d["A2"].font = note

heads = ["코드", "단지명", "배포일", "배포 대수", "문의 수", "구독 전환",
         "문의율", "전환율", "판정", "단지", "회차"]
for i, h in enumerate(heads, start=1):
    c = d.cell(3, i, h)
    c.font, c.fill, c.alignment, c.border = hdr, hdr_fill, ctr, box

for j, code in enumerate(CODES):
    r = DIST_TOP + j
    d.cell(r, 1, code).font = bold
    d.cell(r, 1).alignment = ctr
    for i in (2, 3, 4):                       # 입력 칸
        c = d.cell(r, i)
        c.font, c.fill = inp, in_fill
    d.cell(r, 3).number_format = DATE
    d.cell(r, 4).number_format = NUM
    d.cell(r, 5, f'=COUNTIFS(문의접수!$B${LOG_TOP}:$B${LOG_BOT},$A{r})')
    d.cell(r, 6, f'=COUNTIFS(문의접수!$B${LOG_TOP}:$B${LOG_BOT},$A{r},'
                 f'문의접수!$F${LOG_TOP}:$F${LOG_BOT},"Y")')
    d.cell(r, 7, f'=IFERROR(E{r}/D{r},"")')
    d.cell(r, 8, f'=IFERROR(F{r}/D{r},"")')
    d.cell(r, 9, f'=IF(N(D{r})=0,"미집계",IF(H{r}>=요약!$B$4,"집중","보류"))')
    d.cell(r, 10, f'=MID($A{r},4,1)')
    d.cell(r, 11, f'=MID($A{r},5,1)')
    for i in range(1, 12):
        c = d.cell(r, i)
        c.border = box
        if i >= 5:
            c.font = bod
        if i in (5, 6):
            c.number_format = NUM
        if i in (7, 8):
            c.number_format = PCT
        if i in (9, 10, 11):
            c.alignment = ctr

d.cell(DIST_SUM, 1, "합계").font = bold
for i, col in zip(range(4, 7), "DEF"):
    c = d.cell(DIST_SUM, i, f"=SUM({col}{DIST_TOP}:{col}{DIST_BOT})")
    c.font, c.number_format, c.fill = bold, NUM, sum_fill
d.cell(DIST_SUM, 7, f'=IFERROR(E{DIST_SUM}/D{DIST_SUM},"")')
d.cell(DIST_SUM, 8, f'=IFERROR(F{DIST_SUM}/D{DIST_SUM},"")')
for i in (7, 8):
    c = d.cell(DIST_SUM, i)
    c.font, c.number_format, c.fill = bold, PCT, sum_fill
for i in range(1, 12):
    d.cell(DIST_SUM, i).border = box

d.cell(DIST_SUM + 2, 1,
       "※ '단지'·'회차'는 코드에서 자동으로 뽑은 값입니다. 건드리지 마세요.").font = note
d.cell(DIST_SUM + 3, 1,
       "※ 배포 대수는 배포 당일 차에서 바로 적으세요. "
       "이 숫자가 없으면 전환율의 분모가 없어 실험 전체가 무의미해집니다.").font = note

for col, w in zip("ABCDEFGHIJK", [10, 18, 12, 11, 10, 11, 10, 10, 10, 7, 7]):
    d.column_dimensions[col].width = w
d.freeze_panes = "A4"

# ─────────────────────────────────────────────────────────────
# 3. 문의접수
# ─────────────────────────────────────────────────────────────
g = wb.create_sheet("문의접수")
g.sheet_view.showGridLines = False
g["A1"] = "문의 접수 — 문의 1건당 1행"
g["A1"].font = ttl
g["A2"] = ("문의가 오면 이 시트에 한 줄 추가하세요. QR로 온 문자에는 코드가 이미 찍혀 있고, "
           "전화로 오면 '코드 말씀해주시면 첫 달 1만 원 할인'으로 받아냅니다.")
g["A2"].font = note

lheads = ["접수일", "코드", "이름", "연락처", "유입경로", "구독 전환",
          "첫 방문일", "월 구독료", "메모"]
for i, h in enumerate(lheads, start=1):
    c = g.cell(3, i, h)
    c.font, c.fill, c.alignment, c.border = hdr, hdr_fill, ctr, box

example = ["2026-08-12", "예시", "김○○", "010-0000-0000", "QR문자", "Y",
           "2026-08-14", 80000, "← 예시 행입니다. 코드가 '예시'라 집계에 잡히지 않습니다."]
for i, v in enumerate(example, start=1):
    c = g.cell(LOG_TOP, i, v)
    c.font, c.fill, c.border = note, ex_fill, box
g.cell(LOG_TOP, 8).number_format = WON

for r in range(LOG_TOP + 1, LOG_TOP + 40):
    for i in range(1, 10):
        c = g.cell(r, i)
        c.font, c.border = inp, box
        if i in (1, 7):
            c.number_format = DATE
        if i == 8:
            c.number_format = WON

dv_code = DataValidation(type="list", formula1='"' + ",".join(CODES) + '"',
                         allow_blank=True, showDropDown=False)
dv_route = DataValidation(type="list", formula1='"QR문자,전화,카톡,대면,기타"',
                          allow_blank=True, showDropDown=False)
dv_yn = DataValidation(type="list", formula1='"Y,N"',
                       allow_blank=True, showDropDown=False)
for dv, col in ((dv_code, "B"), (dv_route, "E"), (dv_yn, "F")):
    g.add_data_validation(dv)
    dv.add(f"{col}{LOG_TOP + 1}:{col}{LOG_BOT}")

g.cell(LOG_TOP + 41, 1,
       "※ '구독 전환'에 Y를 넣는 순간 배포기록·요약의 전환율이 함께 움직입니다.").font = note

for col, w in zip("ABCDEFGHI", [12, 10, 12, 16, 11, 11, 12, 12, 40]):
    g.column_dimensions[col].width = w
g.freeze_panes = "A4"

# ─────────────────────────────────────────────────────────────
# 4. 안내
# ─────────────────────────────────────────────────────────────
h = wb.create_sheet("안내")
h.sheet_view.showGridLines = False
h["A1"] = "쓰는 법"
h["A1"].font = ttl

rows = [
    ("", ""),
    ("색 규칙", ""),
    ("노란 칸 / 파란 글씨", "직접 입력하는 곳. 이것만 채우면 됩니다."),
    ("검은 글씨", "수식이 계산합니다. 건드리면 집계가 깨집니다."),
    ("", ""),
    ("순서", ""),
    ("① 배포 당일", "'배포기록'에 단지명·배포일·배포 대수를 적습니다. 차에서 바로."),
    ("② 문의가 올 때", "'문의접수'에 한 줄 추가. 코드는 목록에서 고릅니다."),
    ("③ 구독이 성사되면", "그 행의 '구독 전환'을 Y로 바꿉니다. 나머지는 자동입니다."),
    ("④ 4주 후", "'요약'의 판정을 봅니다. '집중'이 뜬 단지에 자원을 몰아넣습니다."),
    ("", ""),
    ("코드 읽는 법", ""),
    ("WJ-A1", "WJ=운정 · A=단지 · 1=배포 회차"),
    ("", ""),
    ("주의", ""),
    ("배포 대수를 빠뜨리지 마세요", "문의 수는 분자, 배포 대수는 분모입니다. "
                                    "분모가 없으면 전환율이 안 나오고 실험 전체가 무의미해집니다."),
    ("문의접수 4행은 예시입니다", "코드가 '예시'라 집계에 잡히지 않습니다. 지우셔도 됩니다."),
    ("판정 기준 3%", "'요약' B4에서 바꿀 수 있습니다. 플레이북 §7.2 보수 시나리오의 하한값입니다."),
]
for i, (a, b) in enumerate(rows, start=2):
    ca, cb = h.cell(i, 1, a), h.cell(i, 2, b)
    if a and not b:
        ca.font = big
    else:
        ca.font = bold
        cb.font = bod
        cb.alignment = Alignment(wrap_text=True, vertical="top")

h.column_dimensions["A"].width = 26
h.column_dimensions["B"].width = 72

wb.save("배포기록표.xlsx")
print("생성 완료: 배포기록표.xlsx")
